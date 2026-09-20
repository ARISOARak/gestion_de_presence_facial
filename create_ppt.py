from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

# Créer la présentation
prs = Presentation()

# Couleurs sobres
title_color = RGBColor(33, 37, 41)  # Gris foncé
accent_color = RGBColor(52, 73, 94)  # Bleu-gris pour accent
text_color = RGBColor(44, 62, 80)  # Bleu-gris foncé

# Contenu des slides
slides_content = [
    {
        "title": "Python vs R en Data Science",
        "points": [],
        "notes": "Introduction : Cette présentation compare Python et R, les deux langages majeurs en Data Science. Nous allons voir leurs points forts, leurs limites et dans quels cas choisir l’un ou l’autre."
    },
    {
        "title": "Introduction",
        "points": ["Python : polyvalent", "R : statistique", "Objectif : comparaison"],
        "notes": "Python et R dominent la Data Science. Python est polyvalent et orienté production. R est spécialisé en statistiques et analyses de données. L’objectif est de comparer avantages, limites et domaines d’application."
    },
    {
        "title": "Facilité d’utilisation",
        "points": ["Python : simple", "R : plus complexe", "Python = + accessible"],
        "notes": "Python est réputé pour sa syntaxe claire, proche de l’anglais, ce qui le rend idéal pour les débutants. À l’inverse, R a une syntaxe plus particulière qui demande du temps d’adaptation. Conclusion : Python est plus accessible, surtout pour commencer en Data Science."
    },
    {
        "title": "Visualisation de données",
        "points": ["R : ggplot2, lattice", "Python : matplotlib, seaborn", "R reste + fort"],
        "notes": "R est reconnu pour ses bibliothèques de visualisation très puissantes comme ggplot2. Python propose aussi matplotlib, seaborn ou plotly, mais R garde encore un léger avantage en esthétique et richesse des graphiques."
    },
    {
        "title": "Bibliothèques et écosystème",
        "points": ["Python : ML & IA", "R : statistiques avancées", "Choix selon besoin"],
        "notes": "Python possède un vaste écosystème avec NumPy, Pandas, Scikit-learn et TensorFlow, ce qui le rend très fort en Machine Learning et Intelligence Artificielle. R est spécialisé en statistiques, avec des packages comme dplyr ou caret. En résumé, Python est meilleur pour l’IA et la production, R pour l’analyse statistique poussée."
    },
    {
        "title": "Performance et scalabilité",
        "points": ["Python : production", "R : analyses ponctuelles", "Python + adapté gros volumes"],
        "notes": "Python s’intègre bien aux environnements industriels et gère efficacement les gros volumes de données. R est plus adapté aux analyses ponctuelles sur des ensembles de taille moyenne. Conclusion : Python est mieux pour la production et la scalabilité."
    },
    {
        "title": "Communauté et support",
        "points": ["Python : énorme", "R : académique", "Python + universel"],
        "notes": "Python a une communauté mondiale immense, ce qui facilite l’accès aux tutoriels, forums et solutions. R dispose d’une communauté très active dans le domaine académique et de la recherche. Ainsi, Python bénéficie d’un support plus universel."
    },
    {
        "title": "Conclusion générale",
        "points": ["Python : polyvalent", "R : statistiques/visualisation", "Choix selon profil"],
        "notes": "En conclusion, Python est un langage polyvalent, adapté à la production et au Machine Learning. R est très puissant pour l’analyse statistique et la visualisation. Recommandation : Pour les débutants et l’IA, choisir Python. Pour les statisticiens et chercheurs, R reste un excellent choix."
    },
    {
        "title": "Merci / Questions",
        "points": ["Merci pour votre attention !", "Des questions ?"],
        "notes": "Diapositive finale pour remercier le public et inviter aux questions."
    }
]

# Ajouter les slides
for slide_data in slides_content:
    slide_layout = prs.slide_layouts[1]  # Titre + contenu
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    # Titre
    title.text = slide_data["title"]
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = accent_color

    # Contenu
    tf = content.text_frame
    tf.clear()
    for point in slide_data["points"]:
        p = tf.add_paragraph()
        p.text = point
        p.font.size = Pt(22)
        p.font.color.rgb = text_color

    # Notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = slide_data["notes"]

# Sauvegarder
output_path = "Python_vs_R_DataScience_Sobre.pptx"
prs.save(output_path)
print(f"Fichier créé : {output_path}")
